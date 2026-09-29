"""Reading source documents into plain text pages.

Each loader returns a list of ``(page_number, text)`` pairs. Formats with no page concept
return a single page numbered ``None``, which keeps citations honest -- the system never
invents a page number it did not actually read.
"""

from __future__ import annotations

import logging
from collections.abc import Iterable
from pathlib import Path
from typing import Final

logger = logging.getLogger(__name__)

Page = tuple[int | None, str]

MARKDOWN_SUFFIXES: Final = {".md", ".markdown"}
TEXT_SUFFIXES: Final = {".txt", ".text"}

#: Largest total *uncompressed* size accepted from a DOCX.
#:
#: The HTTP layer caps the upload at ~20 MB, but a DOCX is a ZIP archive, so that cap bounds the
#: compressed size only. A 200 KiB file can legitimately declare 200 MiB of contents -- a
#: decompression bomb -- and python-docx would faithfully expand it. The ZIP central directory
#: records each entry's uncompressed size, so the archive can be inspected and rejected *before*
#: a single byte is decompressed.
MAX_DOCX_UNCOMPRESSED_BYTES: Final = 64 * 1024 * 1024

#: Largest compression ratio accepted from a DOCX. Ordinary documents sit far below this; a ratio
#: in the hundreds is the signature of a crafted archive rather than prose.
MAX_DOCX_COMPRESSION_RATIO: Final = 200

#: Largest number of PDF pages read. Bounds work on a file that is small on disk but declares an
#: enormous page tree.
MAX_PDF_PAGES: Final = 2000


class UnsupportedDocumentError(ValueError):
    """Raised when a file's extension has no registered loader."""


class DocumentReadError(RuntimeError):
    """Raised when a file has a known type but cannot be parsed."""


def _load_plaintext(path: Path) -> list[Page]:
    # errors="replace" rather than "strict": a single bad byte in a 300-page manual should
    # degrade one character, not fail the whole ingestion.
    return [(None, path.read_text(encoding="utf-8", errors="replace"))]


def _load_pdf(path: Path) -> list[Page]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise DocumentReadError("pypdf is required to read PDF files") from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise DocumentReadError(f"could not open PDF {path.name}: {exc}") from exc

    if reader.is_encrypted:
        # Attempt the standard empty-password case; refuse anything needing a real password
        # rather than prompting for one.
        try:
            reader.decrypt("")
        except Exception as exc:
            raise DocumentReadError(
                f"{path.name} is password-protected; decrypt it before ingesting"
            ) from exc

    pages: list[Page] = []
    for number, page in enumerate(reader.pages, start=1):
        if number > MAX_PDF_PAGES:
            # Bounds work on a file that is small on disk but declares an enormous page tree.
            # What was read is still usable, so this truncates rather than failing.
            logger.warning(
                "stopping at the %d-page limit for %s; later pages were not read",
                MAX_PDF_PAGES,
                path.name,
            )
            break
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            # One malformed page must not lose the other 299.
            logger.warning("skipping page %d of %s: %s", number, path.name, exc)
            continue
        if text.strip():
            pages.append((number, text))

    if not pages:
        raise DocumentReadError(
            f"{path.name} yielded no extractable text. It is most likely a scanned image "
            "PDF, which needs OCR before it can be ingested."
        )
    return pages


def _guard_docx_archive(path: Path) -> None:
    """Refuse a DOCX whose declared contents are implausibly large before decompressing it.

    Reads only the ZIP central directory, which stores each entry's uncompressed size. Nothing is
    extracted, so a bomb is rejected at negligible cost.

    Raises:
        DocumentReadError: if the file is not a readable ZIP, or declares too much content.
    """
    import zipfile

    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            uncompressed = sum(entry.file_size for entry in entries)
            compressed = sum(entry.compress_size for entry in entries) or 1
    except zipfile.BadZipFile as exc:
        raise DocumentReadError(f"{path.name} is not a valid DOCX (bad ZIP container)") from exc
    except OSError as exc:
        raise DocumentReadError(f"could not read {path.name}: {exc}") from exc

    if uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
        raise DocumentReadError(
            f"{path.name} declares {uncompressed // (1024 * 1024)} MB of uncompressed content, "
            f"above the {MAX_DOCX_UNCOMPRESSED_BYTES // (1024 * 1024)} MB limit; refusing to "
            "expand it"
        )
    ratio = uncompressed / compressed
    if ratio > MAX_DOCX_COMPRESSION_RATIO:
        raise DocumentReadError(
            f"{path.name} has a compression ratio of {ratio:.0f}:1, above the "
            f"{MAX_DOCX_COMPRESSION_RATIO}:1 limit; this is characteristic of a decompression "
            "bomb rather than a document"
        )


def _load_docx(path: Path) -> list[Page]:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise DocumentReadError("python-docx is required to read .docx files") from exc

    _guard_docx_archive(path)

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise DocumentReadError(f"could not open DOCX {path.name}: {exc}") from exc

    blocks: list[str] = [
        paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()
    ]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    if not blocks:
        raise DocumentReadError(f"{path.name} contains no text")
    # DOCX has no stable page model without rendering, so it is one logical page.
    return [(None, "\n\n".join(blocks))]


def load_document(path: Path) -> list[Page]:
    """Read ``path`` into pages of plain text.

    Raises:
        FileNotFoundError: if the path does not exist.
        UnsupportedDocumentError: if the suffix has no loader.
        DocumentReadError: if the file is of a known type but unreadable.
    """
    if not path.is_file():
        raise FileNotFoundError(f"{path} is not a file")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix == ".docx":
        return _load_docx(path)
    if suffix in MARKDOWN_SUFFIXES or suffix in TEXT_SUFFIXES:
        return _load_plaintext(path)
    raise UnsupportedDocumentError(
        f"unsupported file type {suffix!r}; supported: .pdf, .docx, .md, .markdown, .txt"
    )


def iter_documents(root: Path, suffixes: Iterable[str]) -> list[Path]:
    """Return every file under ``root`` whose suffix is in ``suffixes``, sorted for determinism.

    Hidden files and directories are skipped so that ``.git`` and editor scratch files are
    never ingested.
    """
    allowed = {suffix.lower() for suffix in suffixes}
    if root.is_file():
        return [root] if root.suffix.lower() in allowed else []
    found = [
        path
        for path in root.rglob("*")
        if path.is_file()
        and path.suffix.lower() in allowed
        and not any(part.startswith(".") for part in path.relative_to(root).parts)
    ]
    return sorted(found)
